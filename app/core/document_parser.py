"""多格式文档解析：TXT / PDF / Word(docx)。"""

import io
from pathlib import Path

from app.core.logger import logger

SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


def get_file_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def parse_document(file_bytes: bytes, filename: str) -> str:
    """将文件字节解析为纯文本。"""
    ext = get_file_extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"不支持的文件格式: {ext or '(无扩展名)'}，"
            f"当前支持: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    parsers = {
        ".txt": _parse_txt,
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
    }
    text = parsers[ext](file_bytes)
    return _normalize_text(text)


def _normalize_text(text: str) -> str:
    if not text:
        return ""
    lines = [line.strip() for line in text.splitlines()]
    normalized = "\n".join(line for line in lines if line)
    return normalized.strip()


def _parse_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gbk", "gb2312"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文本编码，请使用 UTF-8 或 GBK 编码的 TXT 文件")


def _parse_pdf(file_bytes: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError("请安装 pypdf: pip install pypdf") from exc

    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for i, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(page_text.strip())
        else:
            logger.warning(f"[Parser] PDF 第 {i + 1} 页未提取到文本（可能为扫描件）")

    if not pages:
        raise ValueError("PDF 中未提取到可读文本，扫描版 PDF 需先 OCR")
    return "\n\n".join(pages)


def _parse_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("请安装 python-docx: pip install python-docx") from exc

    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    if not paragraphs:
        raise ValueError("Word 文档中未提取到可读文本")
    return "\n\n".join(paragraphs)
